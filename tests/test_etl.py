"""Pruebas del ETL con un Excel sintético pequeño (no se usan datos reales por
confidencialidad) y una base en memoria con la misma interfaz que
src/db/repository.py. Ubicación: tests/test_etl.py
Ejecutar con: python -m pytest tests
"""
import copy
from contextlib import contextmanager
from datetime import date

import openpyxl
import pandas as pd
import pytest

from src.etl.extract import MESES
from src.etl.pipeline import cargar_incremental, cargar_inicial
from src.etl.validate import ControlError


def hacer_excel(ruta, datos, totales_declarados=None, anio_celda=None):
    """datos: {anio: {clave: (descripcion, [12 valores])}}.
    totales_declarados / anio_celda permiten fabricar archivos con errores."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for anio in sorted(datos, reverse=True):
        ws = wb.create_sheet(f"VTAS {anio}")
        ws.append(["AÑO:", (anio_celda or {}).get(anio, anio)])
        ws.append(["CLAVE", "DESCRIPCION"] + MESES + ["TOTAL"])
        filas = datos[anio]
        for clave, (desc, v) in filas.items():
            ws.append([clave, desc] + list(v) + [sum(x for x in v if x is not None)])
        sumas = [sum((f[1][i] or 0) for f in filas.values()) for i in range(12)]
        n = (totales_declarados or {}).get(anio, len(filas))
        ws.append([f"TOTAL: {n}", None] + sumas + [sum(sumas)])
    wb.save(ruta)


class RepoMemoria:
    def __init__(self):
        self._p = pd.DataFrame(columns=["clave", "descripcion", "historial_completo", "unidad"])
        self._v = pd.DataFrame(columns=["clave", "fecha", "cantidad"])

    @contextmanager
    def transaccion(self, simular=False):
        copia = (self._p.copy(), self._v.copy())
        try:
            yield
        except Exception:
            self._p, self._v = copia
            raise
        if simular:
            self._p, self._v = copia

    def crear_esquema(self):
        pass

    def productos(self):
        return self._p.sort_values("clave").reset_index(drop=True).copy()

    def ultima_fecha(self):
        return None if self._v.empty else max(self._v["fecha"])

    def ventas(self, desde, hasta):
        v = self._v
        return v[(v["fecha"] >= desde) & (v["fecha"] <= hasta)].copy()

    def resumen_ventas(self):
        return len(self._v), float(self._v["cantidad"].sum()) if len(self._v) else 0.0

    def insertar_productos(self, df):
        nuevos = df[~df["clave"].isin(self._p["clave"])]
        self._p = pd.concat([self._p, nuevos], ignore_index=True)
        return len(df)

    def actualizar_descripciones(self, df):
        for r in df.itertuples():
            self._p.loc[self._p["clave"] == r.clave, "descripcion"] = r.descripcion
        return len(df)

    def insertar_ventas(self, df):
        llaves = set(zip(self._v["clave"], self._v["fecha"]))
        nuevas = df[[(c, f) not in llaves for c, f in zip(df["clave"], df["fecha"])]]
        self._v = pd.concat([self._v, nuevas], ignore_index=True)
        return len(df)


@pytest.fixture
def repo():
    return RepoMemoria()


def base_datos():
    """2022-2024. A y B en todos los años (B en kilos); C solo desde 2023; D solo 2024."""
    return {
        2022: {"A": ("PROD A", [5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16]),
               "B": ("GRANEL B", [1.5, 0, 2, 0, 3, 0, 1, 0, 2, 0, 1, 1])},
        2023: {"A": ("PROD A", [4] * 12),
               "B": ("GRANEL B", [1.0] * 12),
               "C": ("PROD C", [2] * 12)},
        2024: {"A": ("PROD A", [3, 3, 3, 3, 3, 3, 3, 3, 3, 3, 0, 0]),
               "B": ("GRANEL B", [2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 0, 0]),
               "C": ("PROD C", [1] * 10 + [0, 0]),
               "D": ("PROD D", [9] * 10 + [0, 0])},
    }


@pytest.fixture
def datos():
    return copy.deepcopy(base_datos())


ANIOS = [2022, 2023, 2024]


def inicial(tmp_path, repo, datos, hasta="2024-03", **kw):
    ruta = tmp_path / "ventas.xlsx"
    hacer_excel(ruta, datos, **kw)
    return ruta, cargar_inicial(ruta, hasta, repo, anios=ANIOS)


# ---------- carga inicial ----------
def test_inicial_productos_y_ventas(tmp_path, repo, datos):
    _, r = inicial(tmp_path, repo, datos)
    p = repo.productos().set_index("clave")
    assert set(p.index) == {"A", "B", "C", "D"}              # vigentes = hoja 2024
    assert p.loc["A", "historial_completo"] == 1 and p.loc["B", "historial_completo"] == 1
    assert p.loc["C", "historial_completo"] == 0 and p.loc["D", "historial_completo"] == 0
    assert p.loc["B", "unidad"] == "KG" and p.loc["A", "unidad"] == "PZA"
    # solo A y B, de 2022-01 a 2024-03 = 27 meses
    assert r.ventas_insertadas == 2 * 27 and repo.resumen_ventas()[0] == 54
    assert set(repo._v["clave"]) == {"A", "B"}
    assert all(f.day == 1 for f in repo._v["fecha"])


def test_inicial_ceros_reales_se_conservan(tmp_path, repo, datos):
    inicial(tmp_path, repo, datos)
    v = repo._v
    cero = v[(v.clave == "B") & (v.fecha == date(2022, 2, 1))]["cantidad"].iloc[0]
    assert cero == 0


def test_inicial_no_se_repite(tmp_path, repo, datos):
    ruta, _ = inicial(tmp_path, repo, datos)
    with pytest.raises(ControlError):
        cargar_inicial(ruta, "2024-03", repo, anios=ANIOS)


def test_simulacion_no_escribe(tmp_path, repo, datos):
    ruta = tmp_path / "v.xlsx"
    hacer_excel(ruta, datos)
    cargar_inicial(ruta, "2024-03", repo, simular=True, anios=ANIOS)
    assert repo.resumen_ventas()[0] == 0 and repo.productos().empty


# ---------- controles: detienen la carga y no escriben ----------
@pytest.mark.parametrize("cambio,control", [
    ("filas", "filas"),
    ("anio", "año"),
    ("negativo", "negativos"),
    ("nulo", "nulos"),
])
def test_controles_detienen(tmp_path, repo, datos, cambio, control):
    kw = {}
    if cambio == "filas":
        kw["totales_declarados"] = {2024: 99}
    if cambio == "anio":
        kw["anio_celda"] = {2023: 2022}
    if cambio == "negativo":
        datos[2023]["A"] = ("PROD A", [-4] + [4] * 11)
    if cambio == "nulo":
        datos[2023]["A"] = ("PROD A", [None] + [4] * 11)
    ruta = tmp_path / "v.xlsx"
    hacer_excel(ruta, datos, **kw)
    with pytest.raises(ControlError) as e:
        cargar_inicial(ruta, "2024-03", repo, anios=ANIOS)
    assert control in {f.control for f in e.value.fallas}
    assert repo.productos().empty and repo.resumen_ventas()[0] == 0


def test_control_duplicados_y_suma(tmp_path, repo, datos):
    import openpyxl
    ruta = tmp_path / "v.xlsx"
    hacer_excel(ruta, datos)
    wb = openpyxl.load_workbook(ruta)
    ws = wb["VTAS 2023"]
    ws["O3"] = 999                      # TOTAL del producto A (fila 3) ya no cuadra con sus meses
    wb.save(ruta)
    with pytest.raises(ControlError) as e:
        cargar_inicial(ruta, "2024-03", repo, anios=ANIOS)
    assert "suma por producto" in {f.control for f in e.value.fallas}
    assert "VTAS 2023" in str(e.value)  # dice dónde falló


def test_mes_sin_ventas_detiene(tmp_path, repo, datos):
    ruta = tmp_path / "v.xlsx"
    hacer_excel(ruta, datos)
    with pytest.raises(ControlError) as e:
        cargar_inicial(ruta, "2024-11", repo, anios=ANIOS)   # nov-2024 suma 0
    assert "mes sin ventas" in {f.control for f in e.value.fallas}


# ---------- carga incremental ----------
def test_incremental_agrega_solo_meses_nuevos(tmp_path, repo, datos):
    ruta, _ = inicial(tmp_path, repo, datos)
    antes = repo._v.copy()
    r = cargar_incremental(ruta, "2024-05", repo)
    assert r.ventas_insertadas == 2 * 2                        # A y B, abril y mayo
    assert repo.resumen_ventas()[0] == 54 + 4
    despues = repo._v
    previos = despues[despues.fecha <= date(2024, 3, 1)].sort_values(["clave", "fecha"]).reset_index(drop=True)
    assert previos.equals(antes.sort_values(["clave", "fecha"]).reset_index(drop=True))


def test_incremental_es_idempotente(tmp_path, repo, datos):
    ruta, _ = inicial(tmp_path, repo, datos)
    cargar_incremental(ruta, "2024-05", repo)
    n = repo.resumen_ventas()
    r = cargar_incremental(ruta, "2024-05", repo)
    assert r.ventas_insertadas == 0 and repo.resumen_ventas() == n


def test_meses_cargados_no_se_modifican(tmp_path, repo, datos):
    inicial(tmp_path, repo, datos)
    datos[2024]["A"] = ("PROD A", [30, 3, 3, 3, 3, 3, 3, 3, 3, 3, 0, 0])   # corrigen enero
    ruta2 = tmp_path / "nuevo.xlsx"
    hacer_excel(ruta2, datos)
    r = cargar_incremental(ruta2, "2024-04", repo)
    enero = repo._v[(repo._v.clave == "A") & (repo._v.fecha == date(2024, 1, 1))]["cantidad"].iloc[0]
    assert enero == 3                                           # no cambió
    assert any("NO se modificaron" in a for a in r.avisos)       # pero se avisa


def test_producto_conocido_ausente_se_carga_en_cero(tmp_path, repo, datos):
    inicial(tmp_path, repo, datos)
    del datos[2024]["B"]
    ruta2 = tmp_path / "nuevo.xlsx"
    hacer_excel(ruta2, datos)
    cargar_incremental(ruta2, "2024-04", repo)
    b = repo._v[(repo._v.clave == "B") & (repo._v.fecha == date(2024, 4, 1))]["cantidad"].iloc[0]
    assert b == 0


def test_producto_nuevo_entra_sin_historial_ni_ventas(tmp_path, repo, datos):
    inicial(tmp_path, repo, datos)
    datos[2024]["E"] = ("PROD E", [0, 0, 0, 5, 5, 0, 0, 0, 0, 0, 0, 0])
    ruta2 = tmp_path / "nuevo.xlsx"
    hacer_excel(ruta2, datos)
    cargar_incremental(ruta2, "2024-05", repo)
    p = repo.productos().set_index("clave")
    assert p.loc["E", "historial_completo"] == 0
    assert "E" not in set(repo._v["clave"])


def test_incremental_unidad_no_cambia(tmp_path, repo, datos):
    inicial(tmp_path, repo, datos)
    datos[2024]["A"] = ("PROD A", [3, 3, 3, 3.5, 3, 3, 3, 3, 3, 3, 0, 0])
    ruta2 = tmp_path / "nuevo.xlsx"
    hacer_excel(ruta2, datos)
    with pytest.raises(ControlError) as e:
        cargar_incremental(ruta2, "2024-04", repo)
    assert "unidad" in {f.control for f in e.value.fallas}


def test_incremental_falla_no_escribe(tmp_path, repo, datos):
    inicial(tmp_path, repo, datos)
    n = repo.resumen_ventas()
    ruta2 = tmp_path / "malo.xlsx"
    hacer_excel(ruta2, datos, totales_declarados={2024: 77})
    with pytest.raises(ControlError):
        cargar_incremental(ruta2, "2024-05", repo)
    assert repo.resumen_ventas() == n


def test_incremental_sin_carga_inicial(tmp_path, repo, datos):
    ruta = tmp_path / "v.xlsx"
    hacer_excel(ruta, datos)
    with pytest.raises(ControlError):
        cargar_incremental(ruta, "2024-05", repo)